import json
import os
import base64
import bcrypt
from datetime import date
from io import BytesIO
from flask import Blueprint, render_template, request, session, redirect, url_for, flash, jsonify, send_file
from docx import Document
from docx.shared import Inches
from auth import role_required
from database import query

secretaria_bp = Blueprint('secretaria', __name__, url_prefix='/secretaria')


def _get_evento():
    id_vista = session.get('id_evento_vista')
    if id_vista:
        evt = query("SELECT * FROM evento WHERE id_evento=%s", (id_vista,), fetch_one=True)
        if evt:
            return evt
    return query("SELECT * FROM evento WHERE es_activo=1 LIMIT 1", fetch_one=True)


def _pct_proyectos_global(id_evento):
    condicion = (
        "p.id_proyecto IS NOT NULL AND "
        "(p.nombre NOT LIKE 'Proyecto Grupo%%' OR "
        " (p.descripcion IS NOT NULL AND TRIM(p.descripcion) != ''))"
    )
    row = query(
        "SELECT COUNT(DISTINCT g.id_grupo) AS total, "
        f"SUM(CASE WHEN {condicion} THEN 1 ELSE 0 END) AS subidos "
        "FROM grupo g LEFT JOIN proyecto p ON p.id_grupo=g.id_grupo "
        "WHERE g.id_evento=%s",
        (id_evento,), fetch_one=True)
    total   = row['total'] or 0
    subidos = int(row['subidos'] or 0)
    return int(round(subidos * 100.0 / total, 0)) if total else 0


def _get_metricas(id_evento):
    participantes = query(
        "SELECT COUNT(*) AS cnt FROM participacion WHERE id_evento=%s",
        (id_evento,), fetch_one=True)['cnt']
    grupos = query(
        "SELECT COUNT(*) AS cnt FROM grupo WHERE id_evento=%s",
        (id_evento,), fetch_one=True)['cnt']
    proyectos = query(
        "SELECT COUNT(*) AS cnt FROM proyecto p JOIN grupo g ON p.id_grupo=g.id_grupo WHERE g.id_evento=%s",
        (id_evento,), fetch_one=True)['cnt']
    pct_row = query(
        "SELECT ROUND(SUM(CASE WHEN estado='Completado' THEN 1 ELSE 0 END)*100.0/COUNT(*),0) AS pct FROM tarea WHERE id_evento=%s",
        (id_evento,), fetch_one=True)
    pct_tareas = int(pct_row['pct'] or 0) if pct_row and pct_row['pct'] is not None else 0
    return {
        'participantes':        participantes,
        'grupos':               grupos,
        'proyectos':            proyectos,
        'pct_proyectos_global': _pct_proyectos_global(id_evento),
        'pct_tareas':           pct_tareas,
    }


@secretaria_bp.route('/dashboard')
@role_required('secretaria')
def dashboard():
    evento = _get_evento()
    if not evento:
        return render_template('secretaria/dashboard.html',
                               evento=None, metricas={},
                               proyectos_recientes=[], tareas_resumen=[],
                               chart_cursos='[]', chart_estados='[]')

    id_evento = evento['id_evento']
    metricas  = _get_metricas(id_evento)

    chart_cursos = query(
        "SELECT c.nombre, COUNT(p.id_proyecto) AS total, c.color "
        "FROM curso c LEFT JOIN grupo g ON c.id_curso=g.id_curso AND g.id_evento=%s "
        "LEFT JOIN proyecto p ON g.id_grupo=p.id_grupo "
        "GROUP BY c.id_curso, c.nombre, c.color ORDER BY total DESC",
        (id_evento,))
    chart_estados = query(
        "SELECT p.estado, COUNT(*) AS total FROM proyecto p "
        "JOIN grupo g ON p.id_grupo=g.id_grupo WHERE g.id_evento=%s GROUP BY p.estado",
        (id_evento,))
    proyectos_recientes = query(
        "SELECT p.id_proyecto, p.nombre AS nombre_proyecto, p.estado, "
        "c.nombre AS nombre_curso, g.id_grupo, e.num_mesa "
        "FROM proyecto p JOIN grupo g ON p.id_grupo=g.id_grupo "
        "JOIN curso c ON g.id_curso=c.id_curso "
        "JOIN espacio e ON g.id_espacio=e.id_espacio "
        "WHERE g.id_evento=%s ORDER BY p.id_proyecto DESC LIMIT 5",
        (id_evento,))
    tareas_resumen = query(
        "SELECT estado, COUNT(*) AS cnt FROM tarea WHERE id_evento=%s GROUP BY estado",
        (id_evento,))

    todas_tareas = query(
        "SELECT t.titulo, t.estado, t.descripcion, t.fecha_limite, t.comentario "
        "FROM tarea t WHERE t.id_evento=%s "
        "ORDER BY FIELD(t.estado,'En proceso','Pendiente','Completado'), t.fecha_limite ASC",
        (id_evento,))

    return render_template('secretaria/dashboard.html',
                           evento=evento, metricas=metricas,
                           proyectos_recientes=proyectos_recientes,
                           tareas_resumen=tareas_resumen,
                           todas_tareas=todas_tareas,
                           chart_cursos=json.dumps(chart_cursos, default=str),
                           chart_estados=json.dumps(chart_estados, default=str))


@secretaria_bp.route('/api/estudiantes')
@role_required('secretaria')
def api_estudiantes():
    evento = _get_evento()
    if not evento:
        return jsonify([])
    rows = query(
        "SELECT pe.nombre, pe.apellido, pe.correo, pe.dni, es.ciclo "
        "FROM participacion pa "
        "JOIN estudiante es ON pa.id_estudiante = es.id_estudiante "
        "JOIN persona pe ON es.id_persona = pe.id_persona "
        "WHERE pa.id_evento = %s "
        "ORDER BY pe.apellido, pe.nombre",
        (evento['id_evento'],))
    return jsonify(rows)


@secretaria_bp.route('/api/grupos')
@role_required('secretaria')
def api_grupos():
    evento = _get_evento()
    if not evento:
        return jsonify([])
    rows = query(
        "SELECT g.id_grupo, c.nombre AS nombre_curso, c.color, e.num_mesa, e.ubicacion, "
        "       p.nombre AS nombre_proyecto, p.estado AS estado_proyecto, "
        "       pe.nombre, pe.apellido, pe.correo, pe.dni, es.ciclo, "
        "       CASE WHEN g.id_lider = es.id_estudiante THEN 1 ELSE 0 END AS es_lider "
        "FROM grupo g "
        "JOIN curso c ON g.id_curso = c.id_curso "
        "JOIN espacio e ON g.id_espacio = e.id_espacio "
        "LEFT JOIN proyecto p ON g.id_grupo = p.id_grupo "
        "JOIN estudiante_grupo eg ON g.id_grupo = eg.id_grupo "
        "JOIN estudiante es ON eg.id_estudiante = es.id_estudiante "
        "JOIN persona pe ON es.id_persona = pe.id_persona "
        "WHERE g.id_evento = %s "
        "ORDER BY g.id_grupo, es_lider DESC, pe.apellido",
        (evento['id_evento'],))
    grupos = {}
    for r in rows:
        gid = r['id_grupo']
        if gid not in grupos:
            grupos[gid] = {
                'id_grupo': gid,
                'nombre_curso': r['nombre_curso'],
                'color': r['color'],
                'num_mesa': r['num_mesa'],
                'ubicacion': r['ubicacion'],
                'nombre_proyecto': r['nombre_proyecto'],
                'estado_proyecto': r['estado_proyecto'],
                'miembros': []
            }
        grupos[gid]['miembros'].append({
            'nombre': r['nombre'],
            'apellido': r['apellido'],
            'correo': r['correo'],
            'dni': r['dni'],
            'ciclo': r['ciclo'],
            'es_lider': bool(r['es_lider'])
        })
    return jsonify(list(grupos.values()))


@secretaria_bp.route('/api/proyectos')
@role_required('secretaria')
def api_proyectos():
    evento = _get_evento()
    if not evento:
        return jsonify([])
    rows = query(
        "SELECT c.id_curso, c.nombre AS nombre_curso, c.color, "
        "       p.id_proyecto, p.nombre AS nombre_proyecto, p.estado, "
        "       p.descripcion, p.tecnologias_usadas, "
        "       g.id_grupo, e.num_mesa, e.ubicacion "
        "FROM proyecto p "
        "JOIN grupo g ON p.id_grupo = g.id_grupo "
        "JOIN curso c ON g.id_curso = c.id_curso "
        "JOIN espacio e ON g.id_espacio = e.id_espacio "
        "WHERE g.id_evento = %s "
        "ORDER BY c.nombre, p.nombre",
        (evento['id_evento'],))
    cursos = {}
    for r in rows:
        cid = r['id_curso']
        if cid not in cursos:
            cursos[cid] = {
                'id_curso': cid,
                'nombre_curso': r['nombre_curso'],
                'color': r['color'],
                'proyectos': []
            }
        cursos[cid]['proyectos'].append({
            'id_proyecto': r['id_proyecto'],
            'nombre_proyecto': r['nombre_proyecto'],
            'estado': r['estado'],
            'descripcion': r['descripcion'],
            'tecnologias_usadas': r['tecnologias_usadas'],
            'num_mesa': r['num_mesa'],
            'ubicacion': r['ubicacion'],
        })
    return jsonify(list(cursos.values()))


@secretaria_bp.route('/api/grupos-sin-proyecto')
@role_required('secretaria')
def api_grupos_sin_proyecto():
    evento = _get_evento()
    if not evento:
        return jsonify([])
    condicion = (
        "p.id_proyecto IS NULL OR "
        "(p.nombre LIKE 'Proyecto Grupo%%' AND "
        " (p.descripcion IS NULL OR TRIM(p.descripcion)=''))"
    )
    rows = query(
        "SELECT g.id_grupo, g.nombre AS nombre_grupo, c.nombre AS nombre_curso, c.color, "
        "e.num_mesa, e.ubicacion, "
        "per.nombre AS est_nombre, per.apellido AS est_apellido, "
        "CASE WHEN g.id_lider=es.id_estudiante THEN 1 ELSE 0 END AS es_lider "
        "FROM grupo g "
        "JOIN curso c ON g.id_curso=c.id_curso "
        "JOIN espacio e ON g.id_espacio=e.id_espacio "
        f"LEFT JOIN proyecto p ON p.id_grupo=g.id_grupo "
        "LEFT JOIN estudiante_grupo eg ON eg.id_grupo=g.id_grupo "
        "LEFT JOIN estudiante es ON eg.id_estudiante=es.id_estudiante "
        "LEFT JOIN persona per ON es.id_persona=per.id_persona "
        f"WHERE g.id_evento=%s AND ({condicion}) "
        "ORDER BY c.nombre, g.id_grupo, es_lider DESC, per.apellido",
        (evento['id_evento'],))
    cursos = {}
    for row in rows:
        cn = row['nombre_curso']
        if cn not in cursos:
            cursos[cn] = {'nombre_curso': cn, 'color': row['color'], 'grupos': {}}
        gi = row['id_grupo']
        if gi not in cursos[cn]['grupos']:
            cursos[cn]['grupos'][gi] = {
                'id_grupo': gi, 'nombre': row['nombre_grupo'],
                'num_mesa': row['num_mesa'], 'ubicacion': row['ubicacion'],
                'integrantes': []
            }
        if row['est_nombre']:
            cursos[cn]['grupos'][gi]['integrantes'].append({
                'nombre': row['est_nombre'], 'apellido': row['est_apellido'],
                'es_lider': bool(row['es_lider'])
            })
    return jsonify([{'nombre_curso': v['nombre_curso'], 'color': v['color'],
                     'grupos': list(v['grupos'].values())} for v in cursos.values()])


@secretaria_bp.route('/tareas')
@role_required('secretaria')
def tareas():
    evento = _get_evento()
    filtro = request.args.get('estado', 'Todos')
    if not evento:
        return render_template('secretaria/tareas.html', tareas=[], evento=None,
                               filtro=filtro, conteos={})

    id_evento = evento['id_evento']
    sql = ("SELECT * FROM tarea WHERE id_evento=%s AND id_secretaria=%s")
    params = [id_evento, session['role_id']]
    if filtro != 'Todos':
        sql += " AND estado=%s"
        params.append(filtro)
    sql += " ORDER BY FIELD(estado,'En proceso','Pendiente','Completado'), fecha_limite ASC"

    lista  = query(sql, params)
    conteos_raw = query(
        "SELECT estado, COUNT(*) AS cnt FROM tarea WHERE id_evento=%s AND id_secretaria=%s GROUP BY estado",
        (id_evento, session['role_id']))
    conteos = {r['estado']: r['cnt'] for r in conteos_raw}
    conteos.setdefault('Pendiente', 0)
    conteos.setdefault('En proceso', 0)
    conteos.setdefault('Completado', 0)
    conteos['Total'] = sum(conteos.values())

    return render_template('secretaria/tareas.html', tareas=lista, evento=evento,
                           filtro=filtro, conteos=conteos)


@secretaria_bp.route('/tarea/nueva', methods=['GET', 'POST'])
@role_required('secretaria')
def tarea_nueva():
    evento = _get_evento()
    error  = None
    if not evento:
        flash('No hay un evento activo.', 'warning')
        return redirect(url_for('secretaria.tareas'))

    if request.method == 'POST':
        titulo           = request.form.get('titulo', '').strip()
        descripcion      = request.form.get('descripcion', '').strip()
        area_involucrada = request.form.get('area_involucrada', '').strip() or None
        fecha_limite     = request.form.get('fecha_limite', '') or None
        if not titulo:
            error = 'El título es obligatorio.'
        else:
            query(
                "INSERT INTO tarea (id_secretaria,id_evento,titulo,area_involucrada,descripcion,estado,comentario,fecha_creacion,fecha_limite) "
                "VALUES(%s,%s,%s,%s,%s,'Pendiente','',CURDATE(),%s)",
                (session['role_id'], evento['id_evento'], titulo, area_involucrada, descripcion, fecha_limite), commit=True)
            flash('Tarea creada correctamente.', 'success')
            return redirect(url_for('secretaria.tareas'))

    return render_template('secretaria/tarea_form.html', evento=evento, error=error, tarea=None)


@secretaria_bp.route('/tarea/<int:id_tarea>/estado', methods=['POST'])
@role_required('secretaria')
def cambiar_estado_tarea(id_tarea):
    nuevo = request.form.get('estado')
    if nuevo not in ('Pendiente', 'En proceso', 'Completado'):
        flash('Estado inválido.', 'danger')
        return redirect(url_for('secretaria.tareas'))

    if nuevo == 'Completado':
        query("UPDATE tarea SET estado=%s, fecha_culminacion=CURDATE() WHERE id_tarea=%s AND id_secretaria=%s",
              (nuevo, id_tarea, session['role_id']), commit=True)
    else:
        query("UPDATE tarea SET estado=%s, fecha_culminacion=NULL WHERE id_tarea=%s AND id_secretaria=%s",
              (nuevo, id_tarea, session['role_id']), commit=True)
    return redirect(url_for('secretaria.tareas'))


@secretaria_bp.route('/tarea/<int:id_tarea>/comentario', methods=['POST'])
@role_required('secretaria')
def agregar_comentario(id_tarea):
    comentario = request.form.get('comentario', '').strip()
    query("UPDATE tarea SET comentario=%s WHERE id_tarea=%s AND id_secretaria=%s",
          (comentario, id_tarea, session['role_id']), commit=True)
    return redirect(url_for('secretaria.tareas'))


@secretaria_bp.route('/asistencia')
@role_required('secretaria')
def asistencia():
    evento = _get_evento()
    if not evento:
        flash('No hay un evento activo.', 'warning')
        return redirect(url_for('secretaria.dashboard'))

    id_evento = evento['id_evento']

    rows = query(
        "SELECT c.id_curso, c.nombre AS nombre_curso, c.color, "
        "       per.nombre, per.apellido, per.dni, per.correo, "
        "       es.id_estudiante, "
        "       COALESCE(pa.asistencia, 0) AS asistencia, "
        "       pa.firma "
        "FROM inscripcion_curso ic "
        "JOIN curso c   ON ic.id_curso      = c.id_curso "
        "JOIN estudiante es ON ic.id_estudiante = es.id_estudiante "
        "JOIN persona per   ON es.id_persona   = per.id_persona "
        "LEFT JOIN participacion pa "
        "       ON pa.id_estudiante = es.id_estudiante AND pa.id_evento = %s "
        "WHERE ic.id_evento = %s "
        "AND c.id_curso IN (SELECT DISTINCT id_curso FROM grupo WHERE id_evento = %s) "
        "ORDER BY c.nombre, per.apellido, per.nombre",
        (id_evento, id_evento, id_evento))

    cursos = {}
    for r in rows:
        cid = r['id_curso']
        if cid not in cursos:
            cursos[cid] = {
                'id_curso':      cid,
                'nombre_curso':  r['nombre_curso'],
                'color':         r['color'],
                'estudiantes':   [],
            }
        cursos[cid]['estudiantes'].append({
            'id_estudiante': r['id_estudiante'],
            'dni':           r['dni'],
            'nombre':        r['nombre'],
            'apellido':      r['apellido'],
            'correo':        r['correo'],
            'asistencia':    bool(r['asistencia']),
            'firma':         r['firma'],
        })

    return render_template('secretaria/asistencia.html',
                           evento=evento,
                           cursos=list(cursos.values()))


@secretaria_bp.route('/asistencia/guardar', methods=['POST'])
@role_required('secretaria')
def asistencia_guardar():
    data          = request.get_json()
    id_estudiante = data.get('id_estudiante')
    id_evento     = data.get('id_evento')
    asistio       = bool(data.get('asistio', False))
    firma         = data.get('firma') or None
    if not asistio:
        firma = None

    query(
        "INSERT INTO participacion (id_evento, id_estudiante, asistencia, firma) "
        "VALUES (%s, %s, %s, %s) "
        "ON DUPLICATE KEY UPDATE asistencia=%s, firma=%s",
        (id_evento, id_estudiante, asistio, firma, asistio, firma),
        commit=True)
    return jsonify({'ok': True})


@secretaria_bp.route('/ajustes', methods=['GET', 'POST'])
@role_required('secretaria')
def ajustes():
    evento = _get_evento()

    if request.method == 'POST':
        edicion     = request.form.get('edicion',     type=int)
        fecha       = request.form.get('fecha',       '').strip()
        hora_inicio = request.form.get('hora_inicio', '').strip()
        hora_fin    = request.form.get('hora_fin',    '').strip()
        semestre    = request.form.get('semestre',    type=int)
        anio        = request.form.get('anio',        type=int)

        if not all([edicion, fecha, hora_inicio, hora_fin, semestre, anio]):
            flash('Todos los campos son obligatorios.', 'danger')
            return redirect(url_for('secretaria.ajustes'))

        if evento:
            query(
                "UPDATE evento SET edicion=%s, fecha=%s, hora_inicio=%s, hora_fin=%s, "
                "semestre=%s, anio=%s WHERE id_evento=%s",
                (edicion, fecha, hora_inicio, hora_fin, semestre, anio, evento['id_evento']),
                commit=True)
        else:
            query(
                "INSERT INTO evento (edicion, fecha, hora_inicio, hora_fin, semestre, anio) "
                "VALUES (%s,%s,%s,%s,%s,%s)",
                (edicion, fecha, hora_inicio, hora_fin, semestre, anio), commit=True)

        flash('Ajustes de ExpoEpics guardados correctamente.', 'success')
        return redirect(url_for('secretaria.ajustes'))

    return render_template('secretaria/ajustes.html', evento=evento)


@secretaria_bp.route('/nuevo-evento', methods=['POST'])
@role_required('secretaria')
def nuevo_evento():
    edicion     = request.form.get('edicion',     type=int)
    fecha       = request.form.get('fecha',       '').strip()
    hora_inicio = request.form.get('hora_inicio', '').strip()
    hora_fin    = request.form.get('hora_fin',    '').strip()
    semestre    = request.form.get('semestre',    type=int)
    anio        = request.form.get('anio',        type=int)

    if not all([edicion, fecha, hora_inicio, hora_fin, semestre, anio]):
        flash('Todos los campos son obligatorios.', 'danger')
        return redirect(url_for('secretaria.ajustes'))

    query("UPDATE evento SET es_activo=0", commit=True)
    id_nuevo = query(
        "INSERT INTO evento (edicion, fecha, hora_inicio, hora_fin, semestre, anio, es_activo) "
        "VALUES (%s,%s,%s,%s,%s,%s,1)",
        (edicion, fecha, hora_inicio, hora_fin, semestre, anio), commit=True)
    session.pop('id_evento_vista', None)

    flash('Nueva ExpoEpics creada correctamente. Ahora es el evento activo.', 'success')
    return redirect(url_for('secretaria.ajustes'))



MESES_ES = ['enero','febrero','marzo','abril','mayo','junio',
            'julio','agosto','septiembre','octubre','noviembre','diciembre']


@secretaria_bp.route('/complementos')
@role_required('secretaria')
def complementos():
    evento = _get_evento()
    if not evento:
        flash('No hay un evento activo.', 'warning')
        return redirect(url_for('secretaria.dashboard'))

    id_evento = evento['id_evento']
    rows = query(
        "SELECT c.id_curso, c.nombre AS nombre_curso, c.color, "
        "       per.nombre, per.apellido, per.dni, "
        "       es.id_estudiante, "
        "       COALESCE(pa.recibio_complemento, 0) AS recibio_complemento "
        "FROM participacion pa "
        "JOIN estudiante es        ON pa.id_estudiante  = es.id_estudiante "
        "JOIN persona per          ON es.id_persona     = per.id_persona "
        "JOIN inscripcion_curso ic ON ic.id_estudiante  = es.id_estudiante "
        "JOIN curso c              ON ic.id_curso       = c.id_curso "
        "WHERE pa.id_evento = %s AND pa.asistencia = 1 "
        "  AND ic.id_evento = %s "
        "  AND c.id_curso IN (SELECT DISTINCT id_curso FROM grupo WHERE id_evento = %s) "
        "ORDER BY c.nombre, per.apellido, per.nombre",
        (id_evento, id_evento, id_evento))

    cursos = {}
    for r in rows:
        cid = r['id_curso']
        if cid not in cursos:
            cursos[cid] = {
                'id_curso':     cid,
                'nombre_curso': r['nombre_curso'],
                'color':        r['color'],
                'estudiantes':  [],
            }
        cursos[cid]['estudiantes'].append({
            'id_estudiante':      r['id_estudiante'],
            'dni':                r['dni'],
            'nombre':             r['nombre'],
            'apellido':           r['apellido'],
            'recibio_complemento': bool(r['recibio_complemento']),
        })

    return render_template('secretaria/complementos.html',
                           evento=evento,
                           cursos=list(cursos.values()))


@secretaria_bp.route('/complementos/marcar', methods=['POST'])
@role_required('secretaria')
def complementos_marcar():
    data          = request.get_json()
    id_estudiante = data.get('id_estudiante')
    id_evento     = data.get('id_evento')
    valor         = int(bool(data.get('valor', 0)))

    query(
        "UPDATE participacion SET recibio_complemento=%s "
        "WHERE id_evento=%s AND id_estudiante=%s",
        (valor, id_evento, id_estudiante), commit=True)
    return jsonify({'ok': True, 'valor': valor})


@secretaria_bp.route('/complementos/marcar-todos', methods=['POST'])
@role_required('secretaria')
def complementos_marcar_todos():
    data      = request.get_json()
    id_curso  = data.get('id_curso')
    id_evento = data.get('id_evento')

    query(
        "UPDATE participacion pa "
        "JOIN inscripcion_curso ic ON ic.id_estudiante = pa.id_estudiante AND ic.id_evento = pa.id_evento "
        "SET pa.recibio_complemento = 1 "
        "WHERE pa.id_evento = %s AND pa.asistencia = 1 AND ic.id_curso = %s",
        (id_evento, id_curso), commit=True)
    return jsonify({'ok': True})


@secretaria_bp.route('/diplomas')
@role_required('secretaria')
def diplomas():
    evento = _get_evento()
    if not evento:
        flash('No hay un evento activo.', 'warning')
        return redirect(url_for('secretaria.dashboard'))

    id_evento = evento['id_evento']
    rows = query(
        "SELECT c.id_curso, c.nombre AS nombre_curso, c.color, "
        "       per.nombre, per.apellido, per.dni, "
        "       es.id_estudiante, COALESCE(pa.certificado, 0) AS certificado "
        "FROM participacion pa "
        "JOIN estudiante es  ON pa.id_estudiante = es.id_estudiante "
        "JOIN persona per    ON es.id_persona    = per.id_persona "
        "JOIN inscripcion_curso ic ON ic.id_estudiante = es.id_estudiante AND ic.id_evento = pa.id_evento "
        "JOIN curso c            ON ic.id_curso = c.id_curso "
        "WHERE pa.id_evento = %s AND pa.asistencia = 1 "
        "  AND c.id_curso IN (SELECT DISTINCT id_curso FROM grupo WHERE id_evento = %s) "
        "ORDER BY c.nombre, per.apellido, per.nombre",
        (id_evento, id_evento))

    cursos = {}
    for r in rows:
        cid = r['id_curso']
        if cid not in cursos:
            cursos[cid] = {
                'id_curso':     cid,
                'nombre_curso': r['nombre_curso'],
                'color':        r['color'],
                'estudiantes':  [],
            }
        cursos[cid]['estudiantes'].append({
            'id_estudiante': r['id_estudiante'],
            'dni':           r['dni'],
            'nombre':        r['nombre'],
            'apellido':      r['apellido'],
            'certificado':   bool(r['certificado']),
        })

    return render_template('secretaria/diplomas.html',
                           evento=evento,
                           cursos=list(cursos.values()))


@secretaria_bp.route('/diplomas/generar', methods=['POST'])
@role_required('secretaria')
def diplomas_generar():
    id_estudiante = request.form.get('id_estudiante', type=int)
    evento = _get_evento()
    if not evento or not id_estudiante:
        flash('Datos inválidos.', 'danger')
        return redirect(url_for('secretaria.diplomas'))

    query(
        "UPDATE participacion SET certificado=1 "
        "WHERE id_evento=%s AND id_estudiante=%s AND asistencia=1",
        (evento['id_evento'], id_estudiante), commit=True)

    return redirect(url_for('secretaria.diplomas_ver', id_estudiante=id_estudiante))


@secretaria_bp.route('/diplomas/ver/<int:id_estudiante>')
@role_required('secretaria')
def diplomas_ver(id_estudiante):
    evento = _get_evento()
    if not evento:
        return redirect(url_for('secretaria.diplomas'))

    info = query(
        "SELECT per.nombre, per.apellido, "
        "       p.nombre AS nombre_proyecto, "
        "       c.nombre AS nombre_curso "
        "FROM participacion pa "
        "JOIN estudiante es  ON pa.id_estudiante = es.id_estudiante "
        "JOIN persona per    ON es.id_persona    = per.id_persona "
        "LEFT JOIN estudiante_grupo eg ON eg.id_estudiante = es.id_estudiante "
        "LEFT JOIN grupo g  ON eg.id_grupo = g.id_grupo AND g.id_evento = pa.id_evento "
        "LEFT JOIN proyecto p  ON g.id_grupo = p.id_grupo "
        "LEFT JOIN curso c     ON g.id_curso = c.id_curso "
        "WHERE pa.id_estudiante=%s AND pa.id_evento=%s AND pa.asistencia=1 "
        "LIMIT 1",
        (id_estudiante, evento['id_evento']), fetch_one=True)

    if not info:
        flash('Estudiante no encontrado o no asistió.', 'danger')
        return redirect(url_for('secretaria.diplomas'))

    fecha_str = None
    if evento.get('fecha'):
        f = evento['fecha']
        fecha_str = f"{f.day} de {MESES_ES[f.month - 1]} de {f.year}"

    doc = query("SELECT firma FROM docente_expoepics WHERE es_director=1 LIMIT 1", fetch_one=True)
    firma_director = doc['firma'] if doc else None

    return render_template('secretaria/diploma_ver.html',
                           info=info,
                           evento=evento,
                           fecha_str=fecha_str,
                           id_estudiante=id_estudiante,
                           firma_director=firma_director)


@secretaria_bp.route('/diplomas/descargar/<int:id_estudiante>')
@role_required('secretaria')
def diplomas_descargar(id_estudiante):
    evento = _get_evento()
    if not evento:
        flash('No hay un evento activo.', 'warning')
        return redirect(url_for('secretaria.diplomas'))

    info = query(
        "SELECT per.nombre, per.apellido, "
        "       p.nombre AS nombre_proyecto, "
        "       c.nombre AS nombre_curso "
        "FROM participacion pa "
        "JOIN estudiante es  ON pa.id_estudiante = es.id_estudiante "
        "JOIN persona per    ON es.id_persona    = per.id_persona "
        "LEFT JOIN estudiante_grupo eg ON eg.id_estudiante = es.id_estudiante "
        "LEFT JOIN grupo g  ON eg.id_grupo = g.id_grupo AND g.id_evento = pa.id_evento "
        "LEFT JOIN proyecto p  ON g.id_grupo = p.id_grupo "
        "LEFT JOIN curso c     ON g.id_curso = c.id_curso "
        "WHERE pa.id_estudiante=%s AND pa.id_evento=%s AND pa.asistencia=1 "
        "LIMIT 1",
        (id_estudiante, evento['id_evento']), fetch_one=True)

    if not info:
        flash('Estudiante no encontrado o no asistió.', 'danger')
        return redirect(url_for('secretaria.diplomas'))

    semestre_str = ''
    if evento.get('anio'):
        sufijo = 'I' if evento.get('semestre') == 1 else 'II'
        semestre_str = f"{evento['anio']}-{sufijo}"

    fecha_str = ''
    if evento.get('fecha'):
        f = evento['fecha']
        fecha_str = f"{f.day} de {MESES_ES[f.month - 1]} de {f.year}"

    reemplazos = {
        'NOMBRE':        f"{info['apellido']}, {info['nombre']}".upper(),
        'EXPO_SEMESTRE': semestre_str,
        'CURSO':         (info['nombre_curso'] or '').upper(),
        'PROYECTO':      (info['nombre_proyecto'] or '').upper(),
        'FECHA':         fecha_str,
    }

    # Firma del director desde la BD
    firma_buf = None
    doc_firma = query(
        "SELECT firma FROM docente_expoepics WHERE es_director=1 LIMIT 1",
        fetch_one=True)
    if doc_firma and doc_firma.get('firma'):
        try:
            img_data = base64.b64decode(doc_firma['firma'].split(',')[1])
            firma_buf = BytesIO(img_data)
        except Exception:
            firma_buf = None

    template_path = os.path.join(
        os.path.dirname(os.path.dirname(__file__)), 'Diploma_ExpoEpics.docx')

    if not os.path.exists(template_path):
        flash('No se encontró la plantilla Diploma_ExpoEpics.docx.', 'danger')
        return redirect(url_for('secretaria.diplomas_ver', id_estudiante=id_estudiante))

    doc = Document(template_path)

    def _reemplazar_parrafo(para):
        for run in para.runs:
            for k, v in reemplazos.items():
                if f'{{{k}}}' in run.text:
                    run.text = run.text.replace(f'{{{k}}}', v)
            if '{FIRMA}' in run.text:
                run.text = ''
                if firma_buf:
                    firma_buf.seek(0)
                    run.add_picture(firma_buf, height=Inches(0.9))

    for para in doc.paragraphs:
        _reemplazar_parrafo(para)
    for tabla in doc.tables:
        for fila in tabla.rows:
            for celda in fila.cells:
                for para in celda.paragraphs:
                    _reemplazar_parrafo(para)

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)

    nombre_archivo = (
        f"Diploma_{info['apellido']}_{info['nombre']}.docx"
        .replace(' ', '_')
    )
    return send_file(
        buf,
        as_attachment=True,
        download_name=nombre_archivo,
        mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    )


@secretaria_bp.route('/cuenta', methods=['GET', 'POST'])
@role_required('secretaria')
def cuenta():
    error = success = None
    if request.method == 'POST':
        actual    = request.form.get('contrasena_actual', '')
        nueva     = request.form.get('nueva_contrasena', '')
        confirmar = request.form.get('confirmar_contrasena', '')
        persona = query("SELECT contrasena FROM persona WHERE id_persona=%s",
                        (session['id_persona'],), fetch_one=True)
        stored  = persona['contrasena'].encode() if isinstance(persona['contrasena'], str) else persona['contrasena']
        if not bcrypt.checkpw(actual.encode(), stored):
            error = 'La contraseña actual es incorrecta.'
        elif len(nueva) < 6:
            error = 'La nueva contraseña debe tener al menos 6 caracteres.'
        elif nueva != confirmar:
            error = 'Las contraseñas no coinciden.'
        else:
            hashed = bcrypt.hashpw(nueva.encode(), bcrypt.gensalt()).decode()
            query("UPDATE persona SET contrasena=%s WHERE id_persona=%s",
                  (hashed, session['id_persona']), commit=True)
            success = 'Contraseña actualizada correctamente.'
    return render_template('secretaria/cuenta.html', error=error, success=success)
